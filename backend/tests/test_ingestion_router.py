"""Unit-level tests for app/services/ingestion/ — the FileCapabilityRegistry
+ DocumentExtractionRouter + adapters. No DB/network required (these run
even without DATABASE_URL, unlike test_ingestion_docx_regression.py's
full-pipeline test).

Covers Part R's test matrix at the router/adapter level: PDF (deferred,
untouched), DOCX, TXT, JPG/PNG/TIFF (deferred to the OCR provider,
untouched), CSV, XLSX, JSON, XML, an unsupported executable-shaped
extension, a corrupt file, a mismatched extension/content pair, and an
encrypted-shaped Office file — plus Part J's cross-format equivalence and
Part K's spreadsheet-must-not-be-assumed-to-be-labs requirement.
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest

from app.services.ingestion.capability_registry import FORMAT_CAPABILITIES, EXPLICITLY_REJECTED_EXTENSIONS
from app.services.ingestion.contract import ExtractionStatus
from app.services.ingestion.router import route_extraction
from tests.fixtures.synthetic_documents import ROMANIAN_DISCHARGE_BILET_DE_IESIRE

ROMANIAN_DISCHARGE_LINES = [line for line in ROMANIAN_DISCHARGE_BILET_DE_IESIRE.strip().splitlines() if line.strip()]


@pytest.fixture
def tmp_file(tmp_path):
    def _make(name: str, data: bytes | str, encoding: str = "utf-8") -> Path:
        path = tmp_path / name
        if isinstance(data, str):
            path.write_text(data, encoding=encoding)
        else:
            path.write_bytes(data)
        return path

    return _make


# ---------------------------------------------------------------------------
# Every registered format actually dispatches to a real handler (never
# "registered but has no handler", REVISION_ADDITIONS-style staleness guard).
# ---------------------------------------------------------------------------


def test_every_registered_extension_has_a_working_dispatch(tmp_file):
    for extension, capability in FORMAT_CAPABILITIES.items():
        path = tmp_file(f"probe{extension}", b"" if capability.magic_bytes else "")
        result = route_extraction(path, path.name)
        assert result.status != ExtractionStatus.UNSUPPORTED_FORMAT or not capability.supported, (
            f"{extension}: supported=True but got UNSUPPORTED_FORMAT ('{result.reason}') — "
            "a registered-but-unhandled extractor kind, not a real format limitation."
        )


# ---------------------------------------------------------------------------
# Part R matrix — one behavior check per format family.
# ---------------------------------------------------------------------------


def test_pdf_and_images_defer_to_the_existing_provider_unchanged(tmp_file):
    pdf = tmp_file("scan.pdf", b"%PDF-1.4\n%%EOF")
    assert route_extraction(pdf, pdf.name).status == ExtractionStatus.DEFER_TO_PROVIDER

    for ext, magic in [(".jpg", b"\xff\xd8\xff"), (".png", b"\x89PNG\r\n\x1a\n"), (".tiff", b"II*\x00")]:
        path = tmp_file(f"scan{ext}", magic + b"\x00" * 16)
        assert route_extraction(path, path.name).status == ExtractionStatus.DEFER_TO_PROVIDER, ext


def test_docx_real_text_extraction(tmp_file):
    import docx

    document = docx.Document()
    for line in ROMANIAN_DISCHARGE_LINES:
        document.add_paragraph(line)
    path = Path(tmp_file("note.docx", b""))  # placeholder, overwritten below
    document.save(path)

    result = route_extraction(path, path.name)
    assert result.status == ExtractionStatus.LOCAL_TEXT
    assert "BILET DE IE" in result.text
    assert "Hidroxiuree" in result.text


def test_txt_bypasses_ocr_entirely(tmp_file):
    path = tmp_file("note.txt", ROMANIAN_DISCHARGE_BILET_DE_IESIRE)
    result = route_extraction(path, path.name)
    assert result.status == ExtractionStatus.LOCAL_TEXT
    assert result.extraction_provider == "local_text"
    assert "BILET DE IE" in result.text


def test_csv_lab_report_parsed_as_table_not_scanned(tmp_file):
    path = tmp_file("labs.csv", "Test,Value,Unit,Range\nHemoglobin,14.2,g/dL,13.5-17.5\nWBC,7.1,10^3/uL,4.0-10.0\n")
    result = route_extraction(path, path.name)
    assert result.status == ExtractionStatus.LOCAL_TEXT
    assert result.tables and result.tables[0]["headers"] == ["Test", "Value", "Unit", "Range"]
    assert len(result.tables[0]["rows"]) == 2


def test_xlsx_lab_report_parsed_as_table(tmp_file, tmp_path):
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Labs"
    sheet.append(["Test", "Value", "Unit"])
    sheet.append(["Hemoglobin", 14.2, "g/dL"])
    path = tmp_path / "labs.xlsx"
    workbook.save(path)

    result = route_extraction(path, path.name)
    assert result.status == ExtractionStatus.LOCAL_TEXT
    assert result.sections == ["Labs"]
    assert result.tables[0]["rows"] == [["Hemoglobin", "14.2", "g/dL"]]


def test_json_detects_fhir_bundle_not_assumed_for_generic_json(tmp_file):
    fhir_bundle = {
        "resourceType": "Bundle",
        "entry": [{"resource": {"resourceType": "Observation", "code": {"text": "Hemoglobin"}, "status": "final"}}],
    }
    fhir_path = tmp_file("export.json", json.dumps(fhir_bundle))
    fhir_result = route_extraction(fhir_path, fhir_path.name)
    assert fhir_result.status == ExtractionStatus.LOCAL_TEXT
    assert fhir_result.structured_payload["detected_format"] == "fhir"

    generic_path = tmp_file("config.json", json.dumps({"theme": "dark", "version": 3}))
    generic_result = route_extraction(generic_path, generic_path.name)
    assert generic_result.status == ExtractionStatus.LOCAL_TEXT
    assert generic_result.structured_payload["detected_format"] == "generic_json"


def test_xml_detects_cda_not_assumed_for_generic_xml(tmp_file):
    cda_xml = (
        '<?xml version="1.0"?><ClinicalDocument xmlns="urn:hl7-org:v3">'
        "<title>Discharge</title><text>Patient improved</text></ClinicalDocument>"
    )
    cda_path = tmp_file("cda.xml", cda_xml)
    cda_result = route_extraction(cda_path, cda_path.name)
    assert cda_result.status == ExtractionStatus.LOCAL_TEXT
    assert cda_result.structured_payload["detected_format"] == "cda"

    generic_path = tmp_file("data.xml", "<catalog><item>widget</item></catalog>")
    generic_result = route_extraction(generic_path, generic_path.name)
    assert generic_result.status == ExtractionStatus.LOCAL_TEXT
    assert generic_result.structured_payload["detected_format"] == "generic_xml"


def test_unsupported_executable_extension_rejected(tmp_file):
    for ext in [".exe", ".dll", ".bat", ".js", ".zip", ".docm"]:
        assert ext in EXPLICITLY_REJECTED_EXTENSIONS
        assert ext not in FORMAT_CAPABILITIES  # never even reaches the router — rejected at the upload door


def test_corrupt_file_extraction_failed(tmp_file):
    path = tmp_file("corrupt.docx", b"this is not a real zip file at all")
    result = route_extraction(path, path.name)
    assert result.status == ExtractionStatus.EXTRACTION_FAILED


def test_mismatched_extension_content_hl7_disguised_as_txt(tmp_file):
    path = tmp_file("weird.txt", "MSH|^~\\&|SENDER|FAC|RECEIVER|FAC2|20260101||ADT^A01|123|P|2.3\r")
    result = route_extraction(path, path.name)
    assert result.status == ExtractionStatus.UNSUPPORTED_FORMAT
    assert "HL7" in result.reason


def test_dicom_disguised_as_pdf_still_detected(tmp_file):
    fake_pdf = b"%PDF-1.4\n" + b"\x00" * (128 - 9) + b"DICM" + b"\x00" * 20
    path = tmp_file("scan.pdf", fake_pdf)
    result = route_extraction(path, path.name)
    assert result.status == ExtractionStatus.UNSUPPORTED_FORMAT
    assert "DICOM" in result.reason


def test_encrypted_office_document_detected(tmp_file):
    ole_header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 40
    path = tmp_file("protected.docx", ole_header)
    result = route_extraction(path, path.name)
    assert result.status == ExtractionStatus.ENCRYPTED


def test_legacy_doc_and_hl7_and_dicom_report_clear_reasons_never_other(tmp_file):
    doc_path = tmp_file("old.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 20)
    doc_result = route_extraction(doc_path, doc_path.name)
    assert doc_result.status == ExtractionStatus.UNSUPPORTED_FORMAT
    assert "re-save" in doc_result.reason.lower()

    hl7_path = tmp_file("message.hl7", "MSH|^~\\&|A|B|C|D|20260101||ADT^A01|1|P|2.3\r")
    hl7_result = route_extraction(hl7_path, hl7_path.name)
    assert hl7_result.status == ExtractionStatus.UNSUPPORTED_FORMAT

    dcm_path = tmp_file("scan.dcm", b"\x00" * 128 + b"DICM" + b"\x00" * 20)
    dcm_result = route_extraction(dcm_path, dcm_path.name)
    assert dcm_result.status == ExtractionStatus.UNSUPPORTED_FORMAT
    assert "diagnostic image interpretation" in dcm_result.reason


# ---------------------------------------------------------------------------
# Part J — cross-format equivalence: the same clinical content in
# different file formats should not change clinical meaning.
# ---------------------------------------------------------------------------


def test_docx_and_txt_of_the_same_content_classify_identically(tmp_file):
    from app.services.document_classifier import classify_document_text

    import docx

    docx_document = docx.Document()
    for line in ROMANIAN_DISCHARGE_LINES:
        docx_document.add_paragraph(line)
    docx_path = tmp_file("discharge.docx", b"")
    docx_document.save(docx_path)

    txt_path = tmp_file("discharge.txt", ROMANIAN_DISCHARGE_BILET_DE_IESIRE)

    docx_result = route_extraction(docx_path, docx_path.name)
    txt_result = route_extraction(txt_path, txt_path.name)
    assert docx_result.status == ExtractionStatus.LOCAL_TEXT
    assert txt_result.status == ExtractionStatus.LOCAL_TEXT

    docx_classification = classify_document_text(docx_result.text)
    txt_classification = classify_document_text(txt_result.text)

    assert docx_classification.document_type == txt_classification.document_type
    assert docx_classification.document_type.value == "discharge_summary"


# ---------------------------------------------------------------------------
# Part K — a lab-shaped spreadsheet classifies as labs; an unrelated one
# must not be falsely treated as lab data.
# ---------------------------------------------------------------------------


def test_lab_spreadsheet_classifies_as_labs_unrelated_spreadsheet_does_not(tmp_file):
    from app.services.document_classifier import classify_document_text

    lab_csv = tmp_file(
        "labs.csv",
        "Test,Value,Unit,Reference Range\n"
        "Hemoglobin,14.8,g/dL,13.5-17.5\n"
        "Leucocite,12.1,10^3/uL,4.0-10.0\n"
        "Creatinina,1.18,mg/dL,0.7-1.3\n"
        "Glicemie,92,mg/dL,70-100\n",
    )
    lab_result = route_extraction(lab_csv, lab_csv.name)
    lab_classification = classify_document_text(lab_result.text)
    assert lab_classification.document_type.value == "laboratory_results"

    unrelated_csv = tmp_file(
        "shopping_list.csv",
        "Item,Quantity,Store\nMilk,2,SuperMarket\nBread,1,Bakery\nEggs,12,SuperMarket\n",
    )
    unrelated_result = route_extraction(unrelated_csv, unrelated_csv.name)
    unrelated_classification = classify_document_text(unrelated_result.text)
    assert unrelated_classification.document_type.value != "laboratory_results"


# ---------------------------------------------------------------------------
# Security (Part M)
# ---------------------------------------------------------------------------


def test_zip_bomb_shaped_docx_refused(tmp_path):
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (250 * 1024 * 1024))
    result = route_extraction(path, path.name)
    assert result.status == ExtractionStatus.EXTRACTION_FAILED
    assert "uncompressed size" in result.reason


def test_macro_bearing_docx_refused(tmp_path):
    path = tmp_path / "macro.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", "<root/>")
        archive.writestr("word/vbaProject.bin", b"fake macro bytes")
    result = route_extraction(path, path.name)
    assert result.status == ExtractionStatus.EXTRACTION_FAILED
    assert "macro" in result.reason.lower()


def test_spreadsheet_formula_read_as_text_never_executed(tmp_path):
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "=1+1"
    sheet["A2"] = '=HYPERLINK("http://example.com","click")'
    path = tmp_path / "formulas.xlsx"
    workbook.save(path)

    result = route_extraction(path, path.name)
    # openpyxl has no formula engine at all — it never evaluates a
    # formula, whether or not a cached value happens to be present (a
    # workbook saved by openpyxl itself, never opened in real Excel, has
    # none). Success here proves extraction completed WITHOUT attempting
    # any evaluation; if the raw formula string surfaces at all it can
    # only be as inert text (still prefixed with "="), never a clicked
    # link or a computed "2".
    assert result.status == ExtractionStatus.LOCAL_TEXT
    assert "2" not in (result.tables[0]["rows"][0] if result.tables and result.tables[0]["rows"] else [])
    for row in result.tables[0]["rows"] if result.tables else []:
        for cell in row:
            if "HYPERLINK" in cell:
                assert cell.startswith("="), "a formula must never surface as anything but inert text"
